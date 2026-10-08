from __future__ import annotations

from pathlib import Path
from typing import Any
from unicodedata import normalize

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None


def load_text_file(file_path: Path) -> str:
    """
    Load plain text file.

    TODO: Add encoding detection, error handling.
    """
    return file_path.read_text(encoding="utf-8")


_PDF_RADICAL_REPLACEMENTS = str.maketrans({"⺠": "民", "⻘": "青", "⻚": "页"})


def _normalize_pdf_text(text: str) -> str:
    return normalize("NFKC", text).translate(_PDF_RADICAL_REPLACEMENTS)


def load_pdf_file(file_path: Path) -> str:
    """
    Load PDF file and extract text content.
    """
    if PdfReader is None:
        raise RuntimeError("PDF support is unavailable because pypdf is not installed")

    try:
        reader = PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            text = page.extract_text(extraction_mode="layout")
            if text:
                pages_text.append(_normalize_pdf_text(text))
        return "\n\n".join(pages_text)
    except Exception as e:
        raise RuntimeError(f"Error reading PDF {file_path.name}: {e}") from e


def load_pdf_parts(file_path: Path) -> list[dict[str, Any]]:
    if PdfReader is None:
        raise RuntimeError("PDF support is unavailable because pypdf is not installed")
    try:
        parts = []
        for page_number, page in enumerate(PdfReader(file_path).pages, 1):
            text = page.extract_text(extraction_mode="layout")
            if text and (normalized := _normalize_pdf_text(text).strip()):
                parts.append({"content": normalized, "metadata": {"page_number": page_number}})
        if not parts:
            raise ValueError(f"No extractable text found in {file_path.name}")
        return parts
    except Exception as exc:
        raise RuntimeError(f"Error reading PDF {file_path.name}: {exc}") from exc


def load_docx_file(file_path: Path) -> str:
    """
    Load DOCX file and extract text content.
    """
    if Document is None:
        raise RuntimeError("DOCX support is unavailable because python-docx is not installed")

    try:
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise RuntimeError(f"Error reading DOCX {file_path.name}: {e}") from e


def load_docx_parts(file_path: Path) -> list[dict[str, Any]]:
    if Document is None:
        raise RuntimeError("DOCX support is unavailable because python-docx is not installed")
    try:
        document = Document(file_path)
        parts: list[dict[str, Any]] = []
        section: str | None = None
        buffer: list[str] = []

        def flush() -> None:
            if buffer:
                parts.append({"content": "\n\n".join(buffer), "metadata": {"section": section} if section else {}})
                buffer.clear()

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = str(getattr(paragraph.style, "name", "") or "")
            if style.lower().startswith("heading"):
                flush()
                section = text
            else:
                buffer.append(text)
        flush()
        if not parts:
            raise ValueError(f"No extractable text found in {file_path.name}")
        return parts
    except Exception as exc:
        raise RuntimeError(f"Error reading DOCX {file_path.name}: {exc}") from exc


def _infer_content_type(file_path: Path) -> str:
    """
    Infer content type from file extension.
    """
    suffix = file_path.suffix.lower()
    content_type_map = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".doc": "application/msword",
    }
    return content_type_map.get(suffix, "application/octet-stream")


def load_file(file_path: Path) -> dict[str, Any]:
    """
    Load a file and extract text content.

    Supports:
      - Text files (.txt)
      - Markdown files (.md, .markdown)
      - PDF files (.pdf) - requires pypdf
      - DOCX files (.docx) - requires python-docx

    Returns dict with:
      - content: str
      - filename: str
      - content_type: str (inferred from extension)
    """
    suffix = file_path.suffix.lower()
    content_type = _infer_content_type(file_path)

    if suffix == ".txt":
        content = load_text_file(file_path)
    elif suffix in [".md", ".markdown"]:
        # Markdown files are read as text but marked with markdown content type
        try:
            content = load_text_file(file_path)
        except Exception as e:
            content = f"[Error reading Markdown {file_path.name}: {e}]"
    elif suffix == ".pdf":
        content = load_pdf_file(file_path)
    elif suffix == ".docx":
        content = load_docx_file(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    if not content.strip():
        raise ValueError(f"No extractable text found in {file_path.name}")

    return {
        "content": content,
        "filename": file_path.name,
        "content_type": content_type,
    }


def load_file_parts(file_path: Path) -> list[dict[str, Any]]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        parts = load_pdf_parts(file_path)
    elif suffix == ".docx":
        parts = load_docx_parts(file_path)
    else:
        loaded = load_file(file_path)
        parts = [{"content": loaded["content"], "metadata": {}}]
    content_type = _infer_content_type(file_path)
    return [
        {
            **part,
            "filename": file_path.name,
            "content_type": content_type,
            "document_id": str(file_path),
        }
        for part in parts
    ]


def load_files(file_paths: list[Path]) -> list[dict[str, Any]]:
    """
    Load multiple files and return list of document dicts.
    """
    documents: list[dict[str, Any]] = []

    for fp in file_paths:
        documents.extend(load_file_parts(fp))

    return documents
